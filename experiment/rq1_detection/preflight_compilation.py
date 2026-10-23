"""Compile the conventional Excel sources without making any LLM calls."""

from __future__ import annotations

import argparse
import hashlib
import json
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from run_conventional_full_pipeline import DEFAULT_WORKBOOK, read_cases, strip_explicit_gold_comments


DEFAULT_OUTPUT = Path(__file__).resolve().parents[2] / "rw_dataset" / "results" / "conventional_full_pipeline" / "compile_preflight.json"
DEFAULT_TRAINING_WORKBOOK = Path(__file__).resolve().parents[2] / "tools" / "solidity_code_with_full_embeddings_final.xlsx"
ROOT = Path(__file__).resolve().parents[2]


def check_ethersolve(bytecode: str) -> int:
    with tempfile.TemporaryDirectory(prefix="convera-cfg-preflight-") as directory:
        folder = Path(directory)
        evm = folder / "contract.evm"
        report = folder / "report.json"
        evm.write_text(bytecode, encoding="ascii")
        result = subprocess.run(
            ["java", "-jar", str(ROOT / "EtherSolve.jar"), "-r", "-j", "-o", str(report), str(evm)],
            cwd=folder, capture_output=True, text=True, timeout=90,
        )
        if result.returncode != 0:
            raise RuntimeError(f"EtherSolve failed: {result.stderr.strip()[:300]}")
        cfg = json.loads(report.read_text(encoding="utf-8"))
        nodes = cfg["runtimeCfg"]["nodes"]
        if not nodes:
            raise RuntimeError("EtherSolve returned an empty runtime CFG")
        return len(nodes)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--training-workbook", type=Path, default=DEFAULT_TRAINING_WORKBOOK)
    parser.add_argument("--check-cfg", action="store_true", help="Also require EtherSolve to build a non-empty runtime CFG")
    args = parser.parse_args()

    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
    from embedding_agent import compile_runtime_bytecode

    training_sheet = load_workbook(args.training_workbook, read_only=True, data_only=True).active
    training_hashes = {
        hashlib.sha256(row[0].encode("utf-8")).hexdigest()
        for row in training_sheet.iter_rows(min_row=2, values_only=True)
        if isinstance(row[0], str)
    }

    results = []
    for case in read_cases(args.workbook):
        source = strip_explicit_gold_comments(case["code"])
        record = {
            "excel_row": case["excel_row"],
            "status": "failed",
            "exact_training_source_overlap": hashlib.sha256(case["code"].encode("utf-8")).hexdigest() in training_hashes,
        }
        try:
            bytecode, contract = compile_runtime_bytecode(source)
            record.update(status="ok", compiled_contract=contract, runtime_hex_chars=len(bytecode))
            if args.check_cfg:
                record["cfg_nodes"] = check_ethersolve(bytecode)
        except Exception as exc:
            record["error"] = f"{type(exc).__name__}: {exc}"
        results.append(record)
        if len(results) % 20 == 0 or len(results) == 1:
            print(f"Compiled {len(results)} cases; failed {sum(row['status'] == 'failed' for row in results)}", flush=True)
    summary = {
        "workbook_sha256": hashlib.sha256(args.workbook.read_bytes()).hexdigest(),
        "source_mode": "explicit_gold_comments_stripped",
        "cfg_checked": args.check_cfg,
        "training_workbook_sha256": hashlib.sha256(args.training_workbook.read_bytes()).hexdigest(),
        "total": len(results),
        "compiled": sum(row["status"] == "ok" for row in results),
        "failed": sum(row["status"] == "failed" for row in results),
        "exact_training_source_overlap": sum(row["exact_training_source_overlap"] for row in results),
        "compiled_overlap": sum(row["status"] == "ok" and row["exact_training_source_overlap"] for row in results),
        "results": results,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"Compilation preflight: {summary['compiled']}/{summary['total']}; details: {args.output}")
    return 0 if not summary["failed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())

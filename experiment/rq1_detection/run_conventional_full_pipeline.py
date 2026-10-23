"""Run every conventional Excel contract through all six CONVERA graph nodes.

Gold labels and historical GPTLens outputs are never passed to the graph.
Each case executes in a separate temporary working directory so the runner
cannot overwrite the project's sample contract or existing pipeline artifacts.
"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import os
import re
import shutil
import sys
import tempfile
import traceback
from datetime import datetime
from pathlib import Path

from dotenv import dotenv_values, load_dotenv
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_WORKBOOK = Path(__file__).resolve().parent / "convera_benchmark_with_results_new.xlsx"
DEFAULT_RESULTS = ROOT / "rw_dataset" / "results" / "conventional_full_pipeline"
DEFAULT_COMPILATION_REPORT = DEFAULT_RESULTS / "compile_preflight.json"
NEO4J_CREDENTIALS = Path(__file__).resolve().parent / "Neo4j-8d51a59e-Created-2026-10-22.txt"
ARTIFACTS = (
    "input.json", "rag_output.json", "parser_output.json",
    "fusion_output_agent.json", "explainer_output.json", "consensus_output.json", "consensus_raw.txt",
    "detected.evm", "report.json", "report_with_rag.json",
)
REQUIRED_PACKAGES = (
    "langgraph", "langchain_openai", "langchain_community", "langchain_chroma",
    "langchain_huggingface", "neo4j", "torch", "torch_geometric",
    "transformers", "solcx",
)


def digest(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def strip_explicit_gold_comments(code: str) -> str:
    """Remove known benchmark annotation comments, preserving Solidity code."""
    def clean_block(match: re.Match[str]) -> str:
        block = match.group(0)
        if "@vulnerable_at_lines" in block or "@source:" in block:
            return "\n" * block.count("\n")
        return block

    code = re.sub(r"/\*.*?\*/", clean_block, code, flags=re.DOTALL)
    return re.sub(r"(?m)^[ \t]*//[ \t]*<(?:yes|no)>[ \t]*<report>[^\n]*", "", code)


def read_cases(workbook: Path) -> list[dict]:
    sheet = load_workbook(workbook, read_only=True, data_only=True).active
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    if headers[:2] != ("label", "code"):
        raise ValueError(f"Expected first Excel columns ('label', 'code'), got {headers[:2]!r}")
    cases = []
    for excel_row, row in enumerate(rows, start=2):
        label, code = row[:2]
        if not label or not isinstance(code, str) or not code.strip():
            raise ValueError(f"Missing label or Solidity source at Excel row {excel_row}")
        cases.append({"excel_row": excel_row, "label": str(label), "code": code})
    return cases


def package_available(name: str) -> bool:
    try:
        return importlib.util.find_spec(name) is not None
    except (ImportError, ModuleNotFoundError):
        return False


def preflight(rag_backend: str = "unsloth") -> list[str]:
    load_dotenv(ROOT / ".env")
    if NEO4J_CREDENTIALS.is_file():
        credentials = dotenv_values(NEO4J_CREDENTIALS)
        for name in ("NEO4J_URI", "NEO4J_USERNAME", "NEO4J_PASSWORD", "NEO4J_DATABASE"):
            if credentials.get(name):
                os.environ[name] = credentials[name]
        os.environ["NEO4J_VECTOR_INDEX"] = "behavior_index"
    os.environ["SEMANTIC_PROVIDER"] = "mistral"
    os.environ["RAG_GENERATION_PROVIDER"] = rag_backend
    os.environ["CONVERA_REQUIRE_REAL_CFG"] = "1"
    os.environ.setdefault("HF_HUB_OFFLINE", "1")
    os.environ.setdefault("TRANSFORMERS_OFFLINE", "1")
    problems = []
    for name in REQUIRED_PACKAGES:
        if not package_available(name):
            problems.append(f"Missing Python package: {name}")
    if rag_backend == "unsloth" and not package_available("unsloth"):
        problems.append("Missing Python package: unsloth (required by original RAG backend)")
    for name in ("NEO4J_URI", "NEO4J_USERNAME", "NEO4J_PASSWORD", "NEO4J_DATABASE", "NEO4J_VECTOR_INDEX"):
        if not os.environ.get(name):
            problems.append(f"Missing configuration: {name}")
    if not (os.environ.get("MISTRAL_API_KEY") or os.environ.get("mistral")):
        problems.append("Missing configuration: MISTRAL_API_KEY (or legacy mistral alias)")
    provider = os.environ.get("LLM_PROVIDER", "").lower().strip()
    if provider == "mistral" or (not provider and (os.environ.get("MISTRAL_API_KEY") or os.environ.get("mistral"))):
        if not (os.environ.get("MISTRAL_API_KEY") or os.environ.get("mistral")):
            problems.append("Missing configuration: MISTRAL_API_KEY")
    elif provider == "openai" and not os.environ.get("OPENAI_API_KEY"):
        problems.append("Missing configuration: OPENAI_API_KEY")
    elif provider not in ("", "ollama", "openai", "mistral"):
        problems.append(f"Unsupported LLM_PROVIDER: {provider}")
    for name in ("java",):
        if not shutil.which(name):
            problems.append(f"Missing executable: {name}")
    for name in ("EtherSolve.jar", "tools/early_fusion.pth", "tools/early_fusion_label_encoder.pkl"):
        if not (ROOT / name).is_file():
            problems.append(f"Missing model/tool asset: {name}")
    return problems


def prepare_sandbox(folder: Path, code: str) -> None:
    (folder / "contracts").mkdir()
    (folder / "tools").mkdir()
    (folder / "contracts" / "sample.sol").write_text(code, encoding="utf-8")
    shutil.copy2(ROOT / "EtherSolve.jar", folder / "EtherSolve.jar")
    for name in ("early_fusion.pth", "early_fusion_label_encoder.pkl"):
        shutil.copy2(ROOT / "tools" / name, folder / "tools" / name)


def run_case(case: dict, case_dir: Path, graph, pipeline, raw_code: bool) -> dict:
    source = case["code"] if raw_code else strip_explicit_gold_comments(case["code"])
    case_dir.mkdir(parents=True, exist_ok=True)
    record = {
        "excel_row": case["excel_row"],
        "gold_label": case["label"],
        "source_sha256": digest(case["code"]),
        "submitted_sha256": digest(source),
        "status": "failed",
    }
    previous_cwd, previous_root = Path.cwd(), pipeline.ROOT
    with tempfile.TemporaryDirectory(prefix="convera-conventional-") as temporary:
        sandbox = Path(temporary)
        prepare_sandbox(sandbox, source)
        try:
            os.chdir(sandbox)
            pipeline.ROOT = sandbox
            with (case_dir / "run.log").open("w", encoding="utf-8", buffering=1) as log:
                # The existing stage tools print large prompts and source code.
                from contextlib import redirect_stderr, redirect_stdout
                with redirect_stdout(log), redirect_stderr(log):
                    state = graph.invoke({})
            record.update({
                "status": "ok",
                "cfg_provenance": "native_solc_ethersolve",
                "rag_vuln_type": state["rag"].get("vuln_type"),
                "fusion_label": state["fusion"].get("Predict"),
                "reason_vuln_type": state["reason"].get("vuln_type"),
                "consensus_decision": state["consensus"].get("decision"),
                "consensus_summary": state["consensus"].get("final_vulnerability_summary"),
            })
        except Exception as exc:
            record["error"] = f"{type(exc).__name__}: {exc}"
            (case_dir / "traceback.txt").write_text(traceback.format_exc(), encoding="utf-8")
        finally:
            os.chdir(previous_cwd)
            pipeline.ROOT = previous_root
            for name in ARTIFACTS:
                artifact = sandbox / name
                if artifact.is_file():
                    shutil.copy2(artifact, case_dir / name)
    (case_dir / "result.json").write_text(json.dumps(record, indent=2, ensure_ascii=False), encoding="utf-8")
    return record


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output-root", type=Path, default=DEFAULT_RESULTS)
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--raw-code", action="store_true", help="Keep source annotations; not recommended for unbiased evaluation")
    parser.add_argument("--rag-backend", choices=("unsloth", "mistral"), default="unsloth")
    parser.add_argument("--preflight-only", action="store_true")
    parser.add_argument("--resume-run", type=Path, help="Resume an existing output directory without repeating successful cases")
    parser.add_argument("--continue-on-error", action="store_true", help="Record isolated failures and continue to later rows")
    parser.add_argument("--compiled-only", action="store_true", help="Run only rows confirmed compilable by the offline preflight")
    parser.add_argument("--compilation-report", type=Path, default=DEFAULT_COMPILATION_REPORT)
    args = parser.parse_args()
    cases = read_cases(args.workbook)
    workbook_sha256 = hashlib.sha256(args.workbook.read_bytes()).hexdigest()
    eligible_cases = cases
    excluded_rows = []
    training_overlap_rows = []
    if args.compiled_only:
        if args.raw_code:
            raise ValueError("--compiled-only requires the preflight's gold-comment-stripped source mode")
        report = json.loads(args.compilation_report.read_text(encoding="utf-8"))
        if report.get("workbook_sha256") != workbook_sha256 or report.get("source_mode") != "explicit_gold_comments_stripped":
            raise ValueError("Compilation report does not match this workbook and source mode")
        if report.get("cfg_checked") is not True:
            raise ValueError("--compiled-only requires a preflight report built with --check-cfg")
        compiled_rows = {row["excel_row"] for row in report["results"] if row["status"] == "ok"}
        if len(report["results"]) != len(cases) or {row["excel_row"] for row in report["results"]} != {case["excel_row"] for case in cases}:
            raise ValueError("Compilation report is missing Excel rows")
        eligible_cases = [case for case in cases if case["excel_row"] in compiled_rows]
        excluded_rows = [case["excel_row"] for case in cases if case["excel_row"] not in compiled_rows]
        training_overlap_rows = [
            row["excel_row"] for row in report["results"]
            if row["status"] == "ok" and row.get("exact_training_source_overlap")
        ]
    print(f"Excel cases: {len(cases)}; compilable: {len(eligible_cases)}; selected: {min(args.limit, len(eligible_cases)) if args.limit else len(eligible_cases)}")
    print("Gold labels are excluded from all model inputs.")
    problems = preflight(args.rag_backend)
    if problems:
        for problem in problems:
            print(f"BLOCKED: {problem}")
        return 2
    if args.preflight_only:
        print("Preflight passed; no model calls made.")
        return 0

    sys.path.insert(0, str(ROOT))
    import langgraph_pipeline as pipeline

    graph = pipeline.build_audit_graph()
    selected = eligible_cases[:args.limit] if args.limit > 0 else eligible_cases
    source_mode = "raw" if args.raw_code else "explicit_gold_comments_stripped"
    if args.resume_run:
        run_dir = args.resume_run.resolve()
        prior = json.loads((run_dir / "progress.json").read_text(encoding="utf-8"))
        expected = {
            "workbook_sha256": workbook_sha256,
            "source_mode": source_mode,
            "rag_generation_provider": args.rag_backend,
            "selected_cases": len(selected),
            "compiled_only": args.compiled_only,
        }
        for field, value in expected.items():
            if prior.get(field) != value:
                raise ValueError(f"Cannot resume: {field} differs from saved run")
        records = {int(row["excel_row"]): row for row in prior["results"]}
    else:
        run_dir = args.output_root / f"run_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        run_dir.mkdir(parents=True, exist_ok=False)
        records = {}
    if args.compiled_only and not (run_dir / "compile_preflight.json").exists():
        shutil.copy2(args.compilation_report, run_dir / "compile_preflight.json")
    chroma_source = ROOT / "knowledge_base" / "chroma_db"
    if chroma_source.is_dir():
        chroma_copy = run_dir / "chroma_db"
        if not chroma_copy.exists():
            shutil.copytree(chroma_source, chroma_copy)
        os.environ["CONVERA_CHROMA_DB"] = str(chroma_copy.resolve())
    consecutive_failures = 0
    for index, case in enumerate(selected, start=1):
        row_number = case["excel_row"]
        if records.get(row_number, {}).get("status") == "ok":
            print(f"[{index}/{len(selected)}] Excel row {row_number}: already complete", flush=True)
            continue
        case_dir = run_dir / f"row_{row_number:04d}"
        if case_dir.exists():
            attempt = 1
            while (run_dir / f"row_{row_number:04d}_retry_{attempt:02d}").exists():
                attempt += 1
            case_dir = run_dir / f"row_{row_number:04d}_retry_{attempt:02d}"
        record = run_case(case, case_dir, graph, pipeline, args.raw_code)
        record["artifact_dir"] = case_dir.name
        record["exact_training_source_overlap"] = row_number in training_overlap_rows
        records[row_number] = record
        ordered_records = [records[row["excel_row"]] for row in selected if row["excel_row"] in records]
        summary = {
            "workbook": str(args.workbook.resolve()),
            "workbook_sha256": workbook_sha256,
            "source_mode": source_mode,
            "semantic_provider": "mistral",
            "rag_generation_provider": args.rag_backend,
            "mistral_model": os.environ.get("MISTRAL_MODEL", "ministral-8b-latest"),
            "cfg_required": True,
            "fusion_feature_dims": {"code": 768, "semantics": 768, "cfg": 64},
            "total_excel_cases": len(cases),
            "compilable_cases": len(eligible_cases),
            "excluded_uncompilable_rows": excluded_rows,
            "exact_training_source_overlap_rows": training_overlap_rows,
            "compiled_only": args.compiled_only,
            "selected_cases": len(selected),
            "completed": sum(row["status"] == "ok" for row in ordered_records),
            "failed": sum(row["status"] == "failed" for row in ordered_records),
            "results": ordered_records,
        }
        (run_dir / "progress.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
        print(f"[{index}/{len(selected)}] Excel row {case['excel_row']}: {record['status']} {record.get('error', '')}", flush=True)
        consecutive_failures = consecutive_failures + 1 if record["status"] == "failed" else 0
        if record["status"] == "failed" and not args.continue_on_error:
            print("Stopping after first failed case; inspect its traceback and log before retrying.")
            break
        if consecutive_failures >= 3:
            print("Stopping after three consecutive failed cases; inspect the backend before resuming.")
            break
    print(f"Results: {run_dir}")
    return 1 if any(row["status"] == "failed" for row in records.values()) else 0


if __name__ == "__main__":
    raise SystemExit(main())
